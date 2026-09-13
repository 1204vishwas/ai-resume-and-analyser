"""Extract raw text from uploaded resume files (PDF / DOCX / TXT)."""
import io


def extract_text(file_stream, filename):
    """Return plain text from a werkzeug FileStorage-like stream.

    Supports .pdf, .docx and .txt. Raises ValueError for unsupported types.
    """
    name = (filename or "").lower()
    data = file_stream.read()

    if name.endswith(".pdf"):
        return _from_pdf(data)
    if name.endswith(".docx"):
        return _from_docx(data)
    if name.endswith(".txt"):
        return _from_txt(data)
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


def _from_pdf(data):
    import pdfplumber
    text_parts = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")
    return "\n".join(text_parts).strip()


def _from_docx(data):
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    # include table cells too
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts).strip()


def _from_txt(data):
    for encoding in ("utf-8", "latin-1"):
        try:
            return data.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore").strip()
